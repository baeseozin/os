# make_ppt.py
# 구현·결과 슬라이드(.pptx) 생성 — 기존 발표 덱 뒤에 붙이는 용도
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

# -----------------------------------
# 색/폰트 (기존 덱의 라이트블루 톤)
# -----------------------------------
BG          = RGBColor(0xF1, 0xF7, 0xFB)
ACCENT      = RGBColor(0x5B, 0xC8, 0xF5)   # 밝은 하늘색 (섹션 번호)
ACCENT_DARK = RGBColor(0x1F, 0x6F, 0xA8)   # 진한 파랑 (코드/값)
TITLE_CLR   = RGBColor(0x3A, 0x44, 0x55)
BODY_CLR    = RGBColor(0x47, 0x55, 0x69)
WHITE       = RGBColor(0xFF, 0xFF, 0xFF)
CARD        = RGBColor(0xFF, 0xFF, 0xFF)
CODE_BG     = RGBColor(0xF3, 0xF1, 0xEC)
FONT        = "맑은 고딕"
MONO        = "Consolas"

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
SW, SH = prs.slide_width, prs.slide_height
BLANK = prs.slide_layouts[6]


def new_slide():
    slide = prs.slides.add_slide(BLANK)
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = BG
    return slide


def textbox(slide, l, t, w, h, lines, size=18, color=BODY_CLR, bold=False,
            font=FONT, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, line_spacing=1.15):
    if isinstance(lines, str):
        lines = [lines]
    tb = slide.shapes.add_textbox(l, t, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.line_spacing = line_spacing
        text, sz, clr, bd = line if isinstance(line, tuple) else (line, size, color, bold)
        r = p.add_run()
        r.text = text
        r.font.size = Pt(sz)
        r.font.color.rgb = clr
        r.font.bold = bd
        r.font.name = font
    return tb


def header(slide):
    # 상단 얇은 라인
    line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, Inches(0.72), SW, Pt(1.4))
    line.fill.solid(); line.fill.fore_color.rgb = ACCENT
    line.line.fill.background(); line.shadow.inherit = False
    # 우상단 점 5개
    x0 = SW - Inches(1.55)
    for i in range(5):
        dot = slide.shapes.add_shape(MSO_SHAPE.OVAL, x0 + Inches(0.26 * i), Inches(0.30),
                                     Inches(0.12), Inches(0.12))
        dot.fill.solid()
        dot.fill.fore_color.rgb = ACCENT if i >= 5 - 1 else (ACCENT_DARK if i >= 3 else RGBColor(0xC9, 0xE6, 0xF5))
        dot.line.fill.background(); dot.shadow.inherit = False


def section_head(slide, number, title):
    header(slide)
    textbox(slide, Inches(0.55), Inches(0.95), Inches(3), Inches(1.0),
            number, size=54, color=ACCENT, bold=True)
    textbox(slide, Inches(0.6), Inches(1.85), Inches(9), Inches(0.7),
            f": {title}", size=26, color=TITLE_CLR, bold=True)


def card(slide, l, t, w, h, fill=CARD):
    shp = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, l, t, w, h)
    shp.fill.solid(); shp.fill.fore_color.rgb = fill
    shp.line.color.rgb = RGBColor(0xD8, 0xE8, 0xF2)
    shp.line.width = Pt(0.75)
    shp.shadow.inherit = False
    try:
        shp.adjustments[0] = 0.06
    except Exception:
        pass
    return shp


# =========================================
# 표지(섹션 구분) 슬라이드
# =========================================
s = new_slide()
header(s)
textbox(s, Inches(0.55), Inches(2.55), Inches(12), Inches(1.0),
        "구현 & 결과", size=48, color=TITLE_CLR, bold=True, align=PP_ALIGN.CENTER)
textbox(s, Inches(0.55), Inches(3.6), Inches(12), Inches(0.7),
        "AI 기반 시스템 오류 자동 해결 마법사 — 구현 내용과 검증", size=20,
        color=BODY_CLR, align=PP_ALIGN.CENTER)


# =========================================
# 06 시스템 아키텍처
# =========================================
s = new_slide()
section_head(s, "06", "시스템 아키텍처")

# 입력 박스
card(s, Inches(0.6), Inches(2.9), Inches(2.5), Inches(1.9), fill=RGBColor(0xE8, 0xF4, 0xFC))
textbox(s, Inches(0.6), Inches(2.95), Inches(2.5), Inches(1.8),
        [("입력", 18, ACCENT_DARK, True),
         ("로그 · 오류코드", 13, BODY_CLR, False),
         ("· 시스템 상태", 13, BODY_CLR, False)],
        align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)

steps = [
    ("1단계  규칙 기반 분류", "RuleClassifier", "키워드·오류코드·시스템 상태로\n4개 유형 1차 분류"),
    ("2단계  유사 사례 검색", "CaseRetriever", "사례 DB에서 점수 기반으로\n비슷한 로그 검색"),
    ("3단계  언어모델 설명 생성", "LlamaDiagnoser", "Llama 3.2 호출 → 원인·해결책을\n함수 호출(JSON)로 생성"),
]
y = Inches(2.9)
x = Inches(3.45)
cw = Inches(3.0)
for title, cls, desc in steps:
    card(s, x, y, cw, Inches(1.9))
    textbox(s, x + Inches(0.15), y + Inches(0.12), cw - Inches(0.3), Inches(1.7),
            [(title, 15, TITLE_CLR, True),
             (cls, 12, ACCENT_DARK, True),
             ("", 6, BODY_CLR, False),
             (desc, 12, BODY_CLR, False)],
            anchor=MSO_ANCHOR.TOP, line_spacing=1.1)
    x = x + cw + Inches(0.12)

# 출력 박스
card(s, Inches(0.6), Inches(5.05), Inches(12.1), Inches(1.6), fill=RGBColor(0xE8, 0xF4, 0xFC))
textbox(s, Inches(0.8), Inches(5.12), Inches(11.7), Inches(1.45),
        [("출력  진단 리포트 (JSON · 함수 호출 report_diagnosis)", 15, ACCENT_DARK, True),
         ("issue_type · confidence · root_cause · evidence · solutions · safe_action_id",
          13, BODY_CLR, False),
         ("유형·신뢰도·safe_action_id는 규칙/사례가 확정 · 원인·해결책 설명은 Llama 3.2가 생성",
          12, BODY_CLR, False)],
        anchor=MSO_ANCHOR.MIDDLE, line_spacing=1.12)


# =========================================
# 07 모듈 구성
# =========================================
s = new_slide()
section_head(s, "07", "모듈 구성")

rows = [
    ("파일", "역할"),
    ("state.py", "DiagnosisState — 입력과 단계별 결과를 담는 상태 객체"),
    ("crew.py", "RuleClassifier · CaseRetriever · LlamaDiagnoser (3단계 핵심 로직)"),
    ("workflow.py", "DiagnosisWorkflow — 입력→분류→검색→진단→리포트 순서 제어"),
    ("main.py", "진입점 — 예시 오류를 넣고 실행 (분리 버전)"),
    ("mmain.py", "위 4개를 합친 단일 파일 버전 (로직 동일)"),
]
tbl_shape = s.shapes.add_table(len(rows), 2, Inches(0.7), Inches(2.7),
                               Inches(11.9), Inches(3.4))
table = tbl_shape.table
table.columns[0].width = Inches(2.7)
table.columns[1].width = Inches(9.2)
for ri, (c0, c1) in enumerate(rows):
    for ci, val in enumerate((c0, c1)):
        cell = table.cell(ri, ci)
        cell.fill.solid()
        cell.fill.fore_color.rgb = ACCENT if ri == 0 else WHITE
        cell.vertical_anchor = MSO_ANCHOR.MIDDLE
        cell.margin_left = Inches(0.15); cell.margin_top = Inches(0.04)
        cell.margin_bottom = Inches(0.04)
        p = cell.text_frame.paragraphs[0]
        r = p.add_run(); r.text = val
        r.font.size = Pt(14 if ri else 15)
        r.font.bold = (ri == 0) or (ci == 0)
        r.font.name = MONO if (ci == 0 and ri != 0) else FONT
        r.font.color.rgb = WHITE if ri == 0 else (ACCENT_DARK if ci == 0 else BODY_CLR)

textbox(s, Inches(0.7), Inches(6.35), Inches(11.9), Inches(0.6),
        "언어모델 Llama 3.2 (ollama · OpenAI 호환) — 접속 정보는 환경변수(LLAMA_BASE_URL · LLAMA_MODEL)로 분리",
        size=13, color=ACCENT_DARK, bold=True)


# =========================================
# 08 동작 예시 (권한 문제)
# =========================================
s = new_slide()
section_head(s, "08", "동작 예시 — 권한 문제")

# 입력
textbox(s, Inches(0.65), Inches(2.7), Inches(6), Inches(0.4),
        "입력 (오류 로그 + 시스템 상태)", size=15, color=TITLE_CLR, bold=True)
card(s, Inches(0.65), Inches(3.1), Inches(5.9), Inches(3.3), fill=CODE_BG)
textbox(s, Inches(0.85), Inches(3.25), Inches(5.5), Inches(3.0),
        "PermissionError: [Errno 13]\n"
        "Access is denied\n"
        "현재 사용자는 관리자 권한이 아님\n\n"
        'system_state = {\n'
        '  "os": "Windows 11",\n'
        '  "is_admin": false\n'
        '}',
        size=13.5, color=ACCENT_DARK, font=MONO, line_spacing=1.25)

# 출력
textbox(s, Inches(6.95), Inches(2.7), Inches(6), Inches(0.4),
        "출력 (함수 호출 report_diagnosis → JSON)", size=15, color=TITLE_CLR, bold=True)
card(s, Inches(6.95), Inches(3.1), Inches(5.75), Inches(3.3), fill=CODE_BG)
textbox(s, Inches(7.15), Inches(3.22), Inches(5.4), Inches(3.1),
        '{\n'
        '  "issue_type": "권한 문제",\n'
        '  "confidence": 0.95,\n'
        '  "root_cause": "쓰기 권한 없음",\n'
        '  "evidence": ["Access is denied", ...],\n'
        '  "solutions": ["관리자 권한으로 재실행", ...],\n'
        '  "safe_action_id": "rerun_as_admin_guide"\n'
        '}',
        size=12.5, color=ACCENT_DARK, font=MONO, line_spacing=1.25)


# =========================================
# 09 검증 결과 (4/4)
# =========================================
s = new_slide()
section_head(s, "09", "검증 결과 — 4개 유형 모두 정확 분류")

rows = [
    ("오류 유형", "분류 정확", "신뢰도", "safe_action_id"),
    ("프로그램 실행 오류", "✓", "0.91", "reinstall_dependency_guide"),
    ("네트워크 연결 오류", "✓", "0.83", "check_network_guide"),
    ("권한 문제", "✓", "0.95", "rerun_as_admin_guide"),
    ("메모리 부족 / 앱 크래시", "✓", "0.95", "free_memory_guide"),
]
tbl_shape = s.shapes.add_table(len(rows), 4, Inches(0.7), Inches(2.8),
                               Inches(11.9), Inches(3.0))
table = tbl_shape.table
table.columns[0].width = Inches(3.4)
table.columns[1].width = Inches(2.0)
table.columns[2].width = Inches(2.0)
table.columns[3].width = Inches(4.5)
for ri, row in enumerate(rows):
    for ci, val in enumerate(row):
        cell = table.cell(ri, ci)
        cell.fill.solid()
        cell.fill.fore_color.rgb = ACCENT if ri == 0 else WHITE
        cell.vertical_anchor = MSO_ANCHOR.MIDDLE
        cell.margin_left = Inches(0.12); cell.margin_top = Inches(0.03)
        cell.margin_bottom = Inches(0.03)
        p = cell.text_frame.paragraphs[0]
        p.alignment = PP_ALIGN.LEFT if ci in (0, 3) else PP_ALIGN.CENTER
        r = p.add_run(); r.text = val
        r.font.size = Pt(14)
        r.font.bold = (ri == 0) or (ci == 0)
        r.font.name = MONO if (ci == 3 and ri != 0) else FONT
        r.font.color.rgb = WHITE if ri == 0 else (ACCENT_DARK if ci in (1, 2) else BODY_CLR)

textbox(s, Inches(0.7), Inches(6.1), Inches(11.9), Inches(0.7),
        "4개 유형 모두 정확 분류  →  4 / 4    ·    신뢰도는 규칙 매칭 강도로 산출",
        size=16, color=ACCENT_DARK, bold=True)


out = "구현_결과_슬라이드.pptx"
prs.save(out)
print(f"saved -> {out}  ({len(prs.slides.__iter__.__self__._sldIdLst)} slides)")
